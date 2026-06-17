
import os
import re
from openpyxl import load_workbook

# ============================================================
# Paths
# ============================================================

SOURCE_DIR = "/data/surendra/HiPC_SRS/ProDo-GNN_CGO/ProDo-GNN_result"
EXCEL_FILE = "/data/surendra/HiPC_SRS/ProDo-GNN_CGO/ProDo-GNN_CGO27.xlsx"

# ============================================================
# Dataset -> ProDo-GNN columns
# ============================================================

DATASET_COLS = {
    "ogbn-arxiv":    ["K",  "L",  "M",  "N",  "O"],
    "reddit":        ["AD", "AE", "AF", "AG", "AH"],
    "ogbn-products": ["AW", "AX", "AY", "AZ", "BA"],
    "igb-small":     ["BP", "BQ", "BR", "BS", "BT"],
    "yelp":          ["CI", "CJ", "CK", "CL", "CM"],
}

# ============================================================
# Batch size -> row offset
# ============================================================

ROW_OFFSET = {
    1024: 0,
    2048: 1,
    4096: 2,
    8192: 3,
    16384: 4,
    32768: 5,
    65536: 6,
}

def get_row(fanout, batch_size):

    if fanout == 10:
        base_row = 8

    elif fanout == 15:
        base_row = 15

    elif fanout == 20:
        base_row = 22

    else:
        return None

    return base_row + ROW_OFFSET[batch_size]

def parse_file(filepath):
    """
    Extract:
        spmm_time
        sampling_time_kernel
        model_time
        total_time
        accuracy
    """

    with open(filepath, "r") as f:
        lines = [line.strip() for line in f.readlines()]

    spmm_time = None
    sampling_kernel = None
    model_time = None
    total_time = None
    accuracy = None

    for line in lines:

        # Example:
        # Sampling time: 56.7697, Model training time: 179.7109, Total time 236.5004

        if "Sampling time:" in line and "Model training time:" in line:

            m = re.search(
                r"Sampling time:\s*([0-9.]+),\s*"
                r"Model training time:\s*([0-9.]+),\s*"
                r"Total time\s*:?\s*([0-9.]+)",
                line
            )

            if m:
                model_time = float(m.group(2))
                total_time = float(m.group(3))

        # Example:
        # Test Accuracy 0.7672

        elif "Test Accuracy" in line:

            m = re.search(r"([0-9]+\.[0-9]+)", line)

            if m:
                accuracy = float(m.group(1))

    # Find:
    #
    # spmm_time, sampling_time
    # 30.096928, 3.365167

    for i in range(len(lines)):

        if (
            "spmm_time" in lines[i]
            and "sampling_time" in lines[i]
            and i + 1 < len(lines)
        ):

            vals = lines[i + 1].split(",")

            if len(vals) >= 2:
                spmm_time = float(vals[0].strip())
                sampling_kernel = float(vals[1].strip())

            break

    return (
        spmm_time,
        sampling_kernel,
        model_time,
        total_time,
        accuracy,
    )


# ============================================================
# Load workbook
# ============================================================

wb = load_workbook(EXCEL_FILE)
ws = wb["main_table"]

updated = 0

for fname in sorted(os.listdir(SOURCE_DIR)):

    if not fname.endswith(".txt"):
        continue

    match = re.search(
        r"(.*)_F(\d+)_B(\d+).*_E(\d+)\.txt",
        fname
    )

    if not match:
        print(f"[SKIP] {fname}")
        continue

    dataset = match.group(1)
    fanout = int(match.group(2))
    batch_size = int(match.group(3))

    if dataset not in DATASET_COLS:
        print(f"[UNKNOWN DATASET] {dataset}")
        continue

    row = get_row(fanout, batch_size)

    if row is None:
        print(f"[INVALID FANOUT] {fname}")
        continue

    filepath = os.path.join(SOURCE_DIR, fname)

    (
        spmm_time,
        sampling_time,
        model_time,
        total_time,
        accuracy,
    ) = parse_file(filepath)

    if None in [
        spmm_time,
        sampling_time,
        model_time,
        total_time,
        accuracy,
    ]:

        print(f"[FAILED] {fname}")
        continue

    cols = DATASET_COLS[dataset]

    ws[f"{cols[0]}{row}"] = spmm_time
    ws[f"{cols[1]}{row}"] = sampling_time
    ws[f"{cols[2]}{row}"] = model_time
    ws[f"{cols[3]}{row}"] = total_time
    ws[f"{cols[4]}{row}"] = accuracy

    updated += 1

    print(
        f"[UPDATED] "
        f"{dataset} "
        f"F={fanout} "
        f"B={batch_size}"
    )

# ============================================================
# Save
# ============================================================

wb.save(EXCEL_FILE)

print()
print("=" * 60)
print(f"Total updated entries : {updated}")
print("Workbook saved successfully.")
print("=" * 60)

