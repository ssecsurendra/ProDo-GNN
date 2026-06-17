```python
import os
import re
from openpyxl import load_workbook

SOURCE_DIR = "/data/surendra/HiPC_SRS/ProDo-GNN_CGO/ProDo-GNN_result"
EXCEL_FILE = "/data/surendra/HiPC_SRS/ProDo-GNN_CGO/ProDo-GNN_CGO27.xlsx"

DATASET_COLS = {
    "ogbn-arxiv":    ["K",  "L",  "M",  "N",  "O"],
    "reddit":        ["AD", "AE", "AF", "AG", "AH"],
    "ogbn-products": ["AW", "AX", "AY", "AZ", "BA"],
    "igb-small":     ["BP", "BQ", "BR", "BS", "BT"],
    "yelp":          ["CI", "CJ", "CK", "CL", "CM"],
}

ROW_MAP = {
    1024: 0,
    2048: 1,
    4096: 2,
    8192: 3,
    16384: 4,
    32768: 5,
    65536: 6,
    131072: 7,
}

def get_row(fanout, batch):
    if fanout == 10:
        base = 8
    elif fanout == 15:
        base = 16
    elif fanout == 20:
        base = 24
    else:
        return None

    return base + ROW_MAP[batch]

wb = load_workbook(EXCEL_FILE)
ws = wb["main_table"]

for fname in os.listdir(SOURCE_DIR):

    if not fname.endswith(".txt"):
        continue

    filepath = os.path.join(SOURCE_DIR, fname)

    m = re.search(
        r"(.*)_F(\d+)_B(\d+).*_E(\d+)\.txt",
        fname
    )

    if not m:
        print("Skipping:", fname)
        continue

    dataset = m.group(1)
    fanout = int(m.group(2))
    batch = int(m.group(3))

    if dataset not in DATASET_COLS:
        print("Unknown dataset:", dataset)
        continue

    row = get_row(fanout, batch)

    with open(filepath, "r") as f:
        lines = [x.strip() for x in f.readlines()]

    model_time = None
    total_time = None
    accuracy = None
    spmm_time = None
    sampling_time = None

    for line in lines:

        if "Model training time:" in line:

            m2 = re.search(
                r"Model training time:\s*([0-9.]+),\s*Total time\s*([0-9.]+)",
                line
            )

            if m2:
                model_time = float(m2.group(1))
                total_time = float(m2.group(2))

        elif line.startswith("Test Accuracy"):
            accuracy = float(line.split()[-1])

    for i in range(len(lines)):

        if "spmm_time,sampling_time" in lines[i]:

            vals = lines[i+1].split(",")

            spmm_time = float(vals[0].strip())
            sampling_time = float(vals[1].strip())
            break

    if None in [spmm_time,
                sampling_time,
                model_time,
                total_time,
                accuracy]:

        print("Failed parsing:", fname)
        continue

    cols = DATASET_COLS[dataset]

    ws[f"{cols[0]}{row}"] = spmm_time
    ws[f"{cols[1]}{row}"] = sampling_time
    ws[f"{cols[2]}{row}"] = model_time
    ws[f"{cols[3]}{row}"] = total_time
    ws[f"{cols[4]}{row}"] = accuracy

    print(
        dataset,
        fanout,
        batch,
        "updated"
    )

wb.save(EXCEL_FILE)

print("Done.")
```

